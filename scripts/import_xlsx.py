#!/usr/bin/env python3
"""Convert "GCSG Meetings and Links.xlsx" into Supabase seed data.

Usage:
    pip install openpyxl
    python3 scripts/import_xlsx.py path/to/GCSG_Meetings_and_Links.xlsx

Writes:
    supabase/seed.sql      – run in the Supabase SQL editor (or `supabase db reset`)
    src/data/seed.json     – offline fallback used when Supabase env vars are missing

Vendors, board members and the curated open action items are not in the
spreadsheet; they live in STATIC below (taken from the original dashboard mockups).
"""
import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent

STATIC = {
    "vendors": [
        {
            "slug": "kindbridge", "name": "Kindbridge", "role": "Helpline, text & chat",
            "status_label": "Live: In Progress", "status_tone": "progress",
            "detail": "Contract signed and countersigned. The 877-GAM-HALT number, live chat and text are being integrated onto the site now.",
            "cost_label": "$90,000 / year", "annual_cost": 90000, "signed": True,
            "risk_label": "Medium Risk", "risk_tone": "progress",
            "risk_points": [
                "$90K annual fee originally non-refundable — being renegotiated to quarterly payments.",
                "Evergreen / auto-renewal clause flagged by counsel (Adam Kaplan); board proceeded anyway — status of the clause itself is unconfirmed.",
                "Brianne has prior board history with Kindbridge — a disclosed conflict of interest the board is tracking.",
            ],
            "sort": 1,
        },
        {
            "slug": "evive", "name": "Evive", "role": "Digital screening platform",
            "status_label": "Signed · Queued", "status_tone": "done",
            "detail": "Contract signed for a 12-month pilot. Integration begins roughly two weeks after Kindbridge goes fully live, so issues can be traced one change at a time.",
            "cost_label": "$25,000 / 12 months", "annual_cost": 25000, "signed": True,
            "risk_label": "Low Risk", "risk_tone": "done",
            "risk_points": [
                "12-month pilot signed, quarterly-friendly terms.",
                "Integration intentionally sequenced after Kindbridge to isolate any issues.",
            ],
            "sort": 2,
        },
        {
            "slug": "gamban", "name": "Gamban", "role": "Gambling-blocking software",
            "status_label": "In Legal Review", "status_tone": "blocked",
            "detail": "Terms set at a self-serve subscription rate. Final document is being converted to an editable format so counsel can redline before signature.",
            "cost_label": "$15,000 / year", "annual_cost": 15000, "signed": False,
            "risk_label": "Unsigned", "risk_tone": "blocked",
            "risk_points": [
                "$15K self-serve subscription rate agreed verbally.",
                "Signature blocked purely on document format (needs editable Word, not PDF/DocuSign).",
            ],
            "sort": 3,
        },
    ],
    "board_members": [
        {"name": "Brianne Doura-Schawohl", "role": "Executive Lead", "initials": "BD", "sort": 1},
        {"name": "Alex Hall", "role": "Board Director", "initials": "AH", "sort": 2},
        {"name": "Matt Smith", "role": "Board Director & Signatory", "initials": "MS", "sort": 3},
        {"name": "Tareq Dowla", "role": "Board Director", "initials": "TD", "sort": 4},
    ],
    "action_items": [
        {"title": "Install Kindbridge chat/text code + swap site number to 877-GAM-HALT", "owner": "Tareq", "status": "overdue", "due_date": "2026-09-28",
         "note": "Target was Sep 28. Live site still shows the old 1-800-MY-RESET number as of Sep 29 — no chat or text widget yet."},
        {"title": "Remove Birches Health / fix SEO confusion on the live site", "owner": "Brianne + Tareq", "status": "not_done", "due_date": None,
         "note": "Birches Health is still listed as a treatment provider on /get-help — the exact thing causing the Kindbridge mix-up."},
        {"title": "Get Gamban contract into an editable format for redlining", "owner": "Brianne / Matt", "status": "in_review", "due_date": None,
         "note": "Blocked on DocuSign/PDF vs. an editable Word doc counsel can mark up."},
        {"title": "Begin Evive integration", "owner": "Brianne + Tareq", "status": "blocked", "due_date": "2026-10-12",
         "note": "Held until ~2 weeks after Kindbridge is fully live, so issues trace to one change at a time. Target ~Oct 12."},
        {"title": "Formalize founding-partner structure (Adam, Tareq, Jonathan, family office)", "owner": "Adam + Tareq", "status": "open", "due_date": None,
         "note": "Needed before any external GCSG raise. No date set."},
        {"title": "Put together a 2026 fundraising plan (tax-year timing)", "owner": "Tareq", "status": "open", "due_date": "2026-12-31",
         "note": "No plan circulated yet in meeting notes or email since the Sep 22 ask."},
        {"title": "Regain & use GCSG website backend access", "owner": "Tareq", "status": "done", "due_date": None,
         "note": "Kim granted Brianne admin access Sep 14; Brianne reset your login Sep 26."},
    ],
}

# Timeline rows that also appear on the public page, keyed by spreadsheet milestone
# prefix -> (public title, public note).
PUBLIC_MILESTONES = {
    "Mission/Vision Drafted": ("Mission Drafted", "Mission, vision and fundraising strategy set."),
    "Vendor Stack Confirmed": ("Vendor Stack Confirmed", "Three-vendor model locked at $130K/year."),
    "Sign Contracts + Launch": ("Contracts Signed", "Kindbridge and Evive executed."),
    "First Official Board Meeting": ("First Board Meeting", "Quorum held; vendor and governance recap."),
    "Kindbridge Kickoff": ("Kindbridge Kickoff", "Go-live MVP defined with Kindbridge’s team."),
    "Kindbridge Go-Live": ("Go-Live Target", "Helpline, text & chat swap — in progress."),
}

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def parse_date(text):
    """Best-effort 'Sep 22, 2026' / '~Oct 12, 2026' / 'Mid-Oct 2026' -> ISO date or None."""
    if not text:
        return None
    t = str(text).lower()
    m = re.search(r"([a-z]{3})[a-z]*\.?\s+(\d{1,2}),?\s+(\d{4})", t)
    if m and m.group(1) in MONTHS:
        return dt.date(int(m.group(3)), MONTHS[m.group(1)], int(m.group(2))).isoformat()
    m = re.search(r"(early|mid|late)?-?\s*([a-z]{3})[a-z]*\s+(\d{4})", t)
    if m and m.group(2) in MONTHS:
        day = {"early": 5, "mid": 15, "late": 25}.get(m.group(1) or "", 1)
        return dt.date(int(m.group(3)), MONTHS[m.group(2)], day).isoformat()
    m = re.search(r"q([1-4])\s+(\d{4})", t)
    if m:
        return dt.date(int(m.group(2)), int(m.group(1)) * 3, 28).isoformat()
    return None


def slugify(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:60]


def clean(v):
    return None if v is None else str(v).strip() or None


def read_resources(ws):
    out, section = [], None
    for row in ws.iter_rows(min_row=3):
        a, b, c = (clean(row[i].value) for i in range(3))
        link = next((cell.hyperlink.target for cell in row[:6] if cell.hyperlink), None)
        if a and not b and not c and a.isupper():
            section = a.title()
            continue
        if a == "Name" or not a:
            continue
        out.append({"section": section, "name": a, "description": b, "kind": c, "url": link,
                    "sort": len(out) + 1})
    return out


def read_meetings(ws):
    meetings, current, last_cat = [], None, None
    header = re.compile(r"^\s*([A-Za-z]+\.? \d{1,2}, \d{4})\s+—\s+(.*)$")
    for row in ws.iter_rows(min_row=3):
        a = row[0].value
        b, d = clean(row[1].value), clean(row[3].value)
        a_s = str(a) if a is not None else ""
        m = header.match(a_s)
        if m:
            date = parse_date(m.group(1))
            parts = [p.strip() for p in m.group(2).split("—", 1)]
            title, subtitle = parts[0], (parts[1] if len(parts) > 1 else None)
            current = {"slug": f"{date}-{slugify(title)}", "date": date, "title": title,
                       "subtitle": subtitle, "source_url": None, "notes": []}
            meetings.append(current)
            last_cat = None
            continue
        if current is None:
            continue
        url = re.search(r"https?://\S+", a_s)
        if (a_s.startswith("Source") or a_s.startswith("→")) and not b:
            if url:
                current["source_url"] = url.group(0)
            continue
        if a_s.strip() == "Category":
            continue
        if not b:
            continue
        cat = a_s.strip()
        if "Decision" in cat:
            kind = "decision"
        elif "Action" in cat:
            kind = "action"
        elif cat == "Discussion":
            kind = "discussion"
        elif re.match(r"^[A-Z][\w +/&()]*:", b):
            kind = "action"  # uncategorised "Owner: do X" rows
        else:
            kind = last_cat or "discussion"
        last_cat = kind
        current["notes"].append({"kind": kind, "body": b, "owner": d,
                                 "sort": len(current["notes"]) + 1})
    return meetings


def status_key(s):
    s = (s or "").lower()
    return {"complete": "complete", "in progress": "in_progress", "upcoming": "upcoming",
            "deferred": "deferred", "overdue": "overdue"}.get(s, "upcoming")


def read_timeline(ws):
    out = []
    for row in ws.iter_rows(min_row=4):
        a, b, c, d = (clean(row[i].value) for i in range(4))
        if not a or not b:
            continue
        pub = next((v for k, v in PUBLIC_MILESTONES.items() if b.startswith(k)), None)
        label = re.sub(r",? \d{4}", "", a).replace(" — Target", "")
        status = status_key(d)
        # The go-live target date has passed without confirmation (per dashboard, Sep 29).
        if b.startswith("Kindbridge Go-Live"):
            status = "overdue"
        out.append({"date_label": label, "sort_date": parse_date(a), "title": b, "detail": c,
                    "status": status, "is_public": pub is not None,
                    "public_title": pub[0] if pub else None,
                    "public_note": pub[1] if pub else None, "sort": len(out) + 1})
    return out


def read_terms(ws):
    out = []
    started = False
    for row in ws.iter_rows():
        a, b, c, d = (clean(row[i].value) for i in range(4))
        if a == "Section":
            started = True
            continue
        if started and a and b:
            out.append({"vendor_slug": "kindbridge", "section": a, "term": b, "detail": c,
                        "risk_level": (d or "").upper(), "sort": len(out) + 1})
    return out


def sql_val(v):
    if v is None:
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return str(v)
    if isinstance(v, list):
        return "array[" + ", ".join(sql_val(x) for x in v) + "]::text[]" if v else "'{}'::text[]"
    return "'" + str(v).replace("'", "''") + "'"


def insert(table, rows, cols):
    if not rows:
        return ""
    vals = ",\n".join("  (" + ", ".join(sql_val(r.get(c)) for c in cols) + ")" for r in rows)
    return f"insert into public.{table} ({', '.join(cols)}) values\n{vals};\n\n"


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(sys.argv[1])
    data = dict(STATIC)
    data["resources"] = read_resources(wb["Resources & Links"])
    data["meetings"] = read_meetings(wb["Meeting Notes"])
    data["milestones"] = read_timeline(wb["Timeline"])
    data["vendor_terms"] = read_terms(wb["Kindbridge Agreement Analysis"])

    (ROOT / "src/data/seed.json").write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n")

    sql = ["-- Generated by scripts/import_xlsx.py — do not edit by hand.\n",
           "truncate public.meeting_notes, public.meetings, public.resources, public.milestones,\n"
           "  public.vendor_terms, public.action_items, public.board_members, public.vendors restart identity cascade;\n\n"]
    sql.append(insert("vendors", data["vendors"], ["slug", "name", "role", "status_label", "status_tone", "detail",
                                                    "cost_label", "annual_cost", "signed", "risk_label", "risk_tone",
                                                    "risk_points", "sort"]))
    sql.append(insert("board_members", data["board_members"], ["name", "role", "initials", "sort"]))
    sql.append(insert("action_items", [dict(a, sort=i + 1) for i, a in enumerate(data["action_items"])],
                      ["title", "owner", "status", "due_date", "note", "sort"]))
    sql.append(insert("milestones", data["milestones"], ["date_label", "sort_date", "title", "detail", "status",
                                                         "is_public", "public_title", "public_note", "sort"]))
    sql.append(insert("resources", data["resources"], ["section", "name", "description", "kind", "url", "sort"]))
    sql.append(insert("vendor_terms", data["vendor_terms"], ["vendor_slug", "section", "term", "detail",
                                                             "risk_level", "sort"]))
    sql.append(insert("meetings", data["meetings"], ["slug", "date", "title", "subtitle", "source_url"]))
    notes = [dict(n, meeting_slug=m["slug"]) for m in data["meetings"] for n in m["notes"]]
    sql.append(insert("meeting_notes", notes, ["meeting_slug", "kind", "body", "owner", "sort"]))
    (ROOT / "supabase/seed.sql").write_text("".join(sql))
    print(f"meetings={len(data['meetings'])} notes={len(notes)} resources={len(data['resources'])} "
          f"milestones={len(data['milestones'])} terms={len(data['vendor_terms'])}")


if __name__ == "__main__":
    main()
